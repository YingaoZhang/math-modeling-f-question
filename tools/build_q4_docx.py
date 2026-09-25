"""Build the Question 4 Word report from its Markdown source.

The builder keeps the report reproducible: tables and figures are read from
the sibling ``outputs/q4`` directory, so no machine-specific paths are stored
in the document.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "q4" / "第四问完整论文.md"
OUTPUT = ROOT / "outputs" / "q4" / "第四问完整论文.docx"


def set_cell_style(cell, header: bool, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "D9D9D9")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(8 if header else 8.5)
            if header:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"(?<!\w)\*([^*]+)\*", r"\1", text)
    return text


def parse_row(line: str) -> list[str]:
    return [re.sub(r"\s+", " ", x.strip()) for x in line.strip().strip("|").split("|")]


def separator(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*", line))


def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(clean_inline(text))
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = len(rows[0])
    # Split very wide tables into readable continuation tables.
    chunks = [list(range(ncols))] if ncols <= 8 else [list(range(0, 6)), list(range(6, ncols))]
    for part, indexes in enumerate(chunks):
        if part:
            add_paragraph(doc, "表格续表")
        table = doc.add_table(rows=len(rows), cols=len(indexes))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        header_trpr = table.rows[0]._tr.get_or_add_trPr()
        header_repeat = OxmlElement("w:tblHeader")
        header_repeat.set(qn("w:val"), "true")
        header_trpr.append(header_repeat)
        for i, row in enumerate(rows):
            for j, source_index in enumerate(indexes):
                cell = table.cell(i, j)
                cell.text = row[source_index] if source_index < len(row) else ""
                set_cell_style(cell, i == 0, "315B78" if i == 0 else ("F3F6F8" if i % 2 == 0 else "FFFFFF"))
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    for name, size in (("Normal", 10.5), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.name = "Microsoft YaHei"
    doc.styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = doc.styles["Title"]._element.find(qn("w:pPr"))
    if title_ppr is None:
        title_ppr = OxmlElement("w:pPr")
        doc.styles["Title"]._element.append(title_ppr)
    title_border = OxmlElement("w:pBdr")
    title_bottom = OxmlElement("w:bottom")
    title_bottom.set(qn("w:val"), "nil")
    title_border.append(title_bottom)
    title_ppr.append(title_border)


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc)
    i = 0
    first = True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if first and line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(line[2:])
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = sub.add_run("数学建模 F 题")
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(90, 90, 90)
            doc.add_page_break()
            first = False
            i += 1
            continue
        if line.startswith("### "):
            add_paragraph(doc, line[4:], "Heading 3")
            i += 1
            continue
        if line.startswith("## "):
            add_paragraph(doc, line[3:], "Heading 1")
            i += 1
            continue
        if line.startswith("# "):
            add_paragraph(doc, line[2:], "Heading 1")
            i += 1
            continue
        if line.startswith("!["):
            m = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
            if m:
                image = SOURCE.parent / m.group(2)
                if image.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(image), width=Inches(6.5))
                    cap = doc.add_paragraph(m.group(1))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_after = Pt(8)
                    for run in cap.runs:
                        run.italic = True
                        run.font.size = Pt(9)
                else:
                    add_paragraph(doc, m.group(1))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and separator(lines[i + 1]):
            rows = [parse_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        if re.match(r"^\s*[-*] ", line):
            add_paragraph(doc, re.sub(r"^\s*[-*] ", "", line), "List Bullet")
            i += 1
            continue
        if re.match(r"^\s*\d+\. ", line):
            add_paragraph(doc, re.sub(r"^\s*\d+\. ", "", line), "List Number")
            i += 1
            continue
        add_paragraph(doc, line)
        i += 1
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("问题四  模型效率前沿与规模技术进步分解")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(120, 120, 120)
    doc.core_properties.title = "问题四 模型效率前沿与规模技术进步分解"
    doc.core_properties.subject = "数学建模 F 题第四问"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
