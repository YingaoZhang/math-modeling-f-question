"""Build a polished Word version of the Question 3 Markdown report."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "q3" / "第三问完整论文.md"
OUTPUT = ROOT / "outputs" / "q3" / "第三问完整论文.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, header: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8 if header else 8.5)
    if header:
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def parse_table_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [re.sub(r"\s+", " ", x.strip()) for x in line.split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*", line))


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    text = text.replace("`", "")
    text = re.sub(r"(?<!\w)\*([^*]+)\*", r"\1", text)
    return text


def add_body_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(clean_inline(text))
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def add_equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    equation = text.strip().strip("$").strip()
    if equation.startswith("C_{train}"):
        equation = (
            "C_train = 0.006 N D;   C_attn = C_train · ηL_ctx / 6;   "
            "C_Q = 10⁻¹² D [g(Q) − g(Q₀)]₊"
        )
    elif equation.startswith("L=A N"):
        equation = (
            "L = A N⁻ᵅ + B D⁻ᵝ + γ(1 − Q_eff)ᶿ "
            "+ 0.08 tᵀ(p − p₀) + 0.03 Σᵢ pᵢ ln(pᵢ / p₀ᵢ)"
        )
    run = p.add_run(equation)
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    run.font.size = Pt(10.5)
    run.italic = True


def add_table(doc: Document, rows: list[list[str]], title: str | None = None) -> None:
    if not rows:
        return
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9.5)
    ncols = len(rows[0])
    is_wide = ncols > 8
    indexes = list(range(min(ncols, 6)))
    display_rows = rows if not is_wide else [rows[0], *rows[1:9]]
    table = doc.add_table(rows=len(display_rows), cols=len(indexes))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(display_rows):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if i == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            tr_pr.append(repeat_header)
        for j, source_index in enumerate(indexes):
            cell = table.cell(i, j)
            set_cell_text(cell, row[source_index] if source_index < len(row) else "", header=i == 0)
            set_cell_borders(cell)
            set_cell_shading(cell, "3F6B85" if i == 0 else ("F4F7F9" if i % 2 == 0 else "FFFFFF"))
    if is_wide:
        note = doc.add_paragraph("宽表正文仅列前六个识别字段及代表性记录；完整逐项结果见 outputs/q3 下对应 CSV。")
        note.paragraph_format.space_before = Pt(2)
        note.paragraph_format.space_after = Pt(4)
        for run in note.runs:
            run.italic = True
            run.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size, color in [("Heading 1", 15, "000000"), ("Heading 2", 12.5, "000000"), ("Heading 3", 11.5, "000000")]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
    title = doc.styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if first_heading and line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(line[2:].replace("：", " "))
            p.paragraph_format.space_after = Pt(14)
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p2.add_run("数学建模 F 题")
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(90, 90, 90)
            doc.add_page_break()
            first_heading = False
            i += 1
            continue
        if line.startswith("### "):
            add_body_paragraph(doc, line[4:], "Heading 3")
            i += 1
            continue
        if line.startswith("## "):
            add_body_paragraph(doc, line[3:], "Heading 1")
            i += 1
            continue
        if line.startswith("# "):
            add_body_paragraph(doc, line[2:], "Heading 1")
            i += 1
            continue
        if line.startswith("$$") and line.endswith("$$") and len(line) > 4:
            add_equation(doc, line[2:-2])
            i += 1
            continue
        if line.startswith("!["):
            m = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
            if m:
                image_path = SOURCE.parent / m.group(2)
                if image_path.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(5)
                    p.paragraph_format.space_after = Pt(3)
                    p.add_run().add_picture(str(image_path), width=Inches(6.5))
                    cap = doc.add_paragraph(m.group(1))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_after = Pt(8)
                    cap.runs[0].italic = True
                    cap.runs[0].font.size = Pt(9)
                else:
                    add_body_paragraph(doc, m.group(1))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            table_lines = [parse_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, table_lines)
            continue
        if re.match(r"^\s*[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            p.add_run(clean_inline(re.sub(r"^\s*[-*] ", "", line)))
            i += 1
            continue
        if re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            p.add_run(clean_inline(re.sub(r"^\s*\d+\. ", "", line)))
            i += 1
            continue
        add_body_paragraph(doc, line)
        i += 1
    # Footer with a neutral document label; no page fields are needed for the requested format.
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("问题三  算力约束下的多维资源联合优化")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(120, 120, 120)
    doc.core_properties.title = "问题三 算力约束下的多维资源联合优化"
    doc.core_properties.subject = "数学建模 F 题第三问"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
