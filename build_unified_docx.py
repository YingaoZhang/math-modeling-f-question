"""Build the single integrated paper for Questions 1 and 2."""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
Q1 = ROOT / "outputs" / "q1"
Q2 = ROOT / "outputs" / "q2"
DOC = ROOT / "outputs" / "数学建模F题统一论文.docx"


def borders(table):
    tbl_pr = table._tbl.tblPr
    border = tbl_pr.first_child_found_in("w:tblBorders")
    if border is None:
        border = OxmlElement("w:tblBorders")
        tbl_pr.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = border.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            border.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def paragraph(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    return p


def add_table(doc, frame, digits=4, max_rows=18):
    if frame is None or len(frame) == 0:
        return
    data = frame.head(max_rows).copy()
    table = doc.add_table(rows=1, cols=len(data.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table)
    for j, col in enumerate(data.columns):
        cell = table.rows[0].cells[j]
        cell.text = str(col)
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(7.2)
    for _, row in data.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(data.columns):
            value = row[col]
            if isinstance(value, float):
                text = "" if pd.isna(value) else f"{value:.{digits}f}"
            else:
                text = str(value)
            cells[j].text = text
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7.2)
        if len(table.rows) % 2 == 0:
            for cell in cells:
                shade(cell, "F3F6F9")


def add_figure(doc, path, caption, width=6.15):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in c.runs:
        run.italic = True
        run.font.size = Pt(9)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65)
    sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.75)
    sec.right_margin = Inches(.75)
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Microsoft YaHei"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("数学建模 F 题统一建模论文").bold = True
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("语料质量评分与大模型训练标度律").italic = True
    paragraph(doc, "本文将问题一的语料质量评分、成分配比建模与问题二的参数规模标度律放在同一数据边界下讨论。核心结论是：质量指标总体同向，极端冲突只占少数；The Pile Loss 与 Pythia val_loss 保持分开；B1 采用无截距标度律，Q_score 按噪声缺陷率进入非对称质量项，配比接口暂作为可识别性明确的下一步模型。")

    doc.add_heading("摘要", 1)
    paragraph(doc, "第一问以 A1 全局 ECDF 将 22 个质量信号映射到 (0,1]，按教育、可读性、推理、清洁和结构五组等权构造文档质量分 Q，并以 CRITIC、熵权和长度加权作敏感性分析。A1 的逐域独立置换零分布均值为约 55.7%，实测 A1、A2、A3 的域等权 P75/P25 冲突率为 42.6%、53.7%、47.9%，均处于零分布下侧；P90/P10 主口径下 A1 约 9.1%。因此质量维度总体同向，冲突只属少数情形。配比部分闭合到单纯形并用零值乘性替换和 Helmert ILR，13 个 Loss 域独立建模；线性 Ridge、二次交互和指数 Data Mixing Laws 的宏平均 R² 分别为 0.7641、0.8628 和 0.8688，CV 选型宏平均 R² 为 0.8896。")
    paragraph(doc, "第二问以 B1 的完整 N-D 网格为主，采用 L=A N^-alpha+B D^-beta，先固定 beta=0.2799 再加先验惩罚局部松弛。B6--B8 中 Q_score 与 Loss 同向，故用 L=A N^-alpha+B D^-beta+gamma Q_score^theta，gamma、theta>0；这保证质量提升时损失下降，并在质量趋于完美时退化为经典标度律。跨语料仿射桥接只作为可检验假设，High 层配对仅 7/75 条，当前不可稳定识别。")

    doc.add_heading("1 问题重述与数据边界", 1)
    paragraph(doc, "问题一要求从多维文档质量信号得到域级质量分，识别指标冲突，并解释不同语料配比对 The Pile 验证 Loss 的影响。问题二要求在 Pythia/Cerebras 等 B 表上建立参数量 N、数据量 D 与验证 Loss 的标度关系，并检验质量变量和配比接口的边际意义。A 表 Loss 是 The Pile 验证集交叉熵，B 表 val_loss 是 Pythia 自有验证集，二者不直接混算。A12--A15 是估算数据，A2/A3 含有 A1 的子样本，B9/B10 是超百亿规模情景外推。")
    paragraph(doc, "A1 含 51230 条记录、7 个文档域；问题一的 Loss 配方为 17 个组成域，其中只有 13 个域有对应 Loss。当前 Qbar 接口只有 3 个域有文档级质量观测，14 个域采用总体中位数插补，故 Qbar 作为弱证据的场景量，不被解释为已识别的因果质量变量。")

    doc.add_heading("2 假设与符号", 1)
    add_table(doc, pd.DataFrame([
        ["p", "17 维配比，闭合后满足 sum p_i=1"],
        ["z(p)", "16 维 Helmert ILR 坐标，零值先乘性替换"],
        ["Q", "A 表文档质量分，严格落在 (0,1]"],
        ["Q_score", "B6--B8 原生噪声缺陷率，越大表示缺陷越多"],
        ["N,D", "参数量和训练数据量，单位沿用 B 表"],
        ["L", "逐域验证 Loss；不对 13 个原始 Loss 取算术平均作为因变量"],
    ], columns=["符号", "定义"]), digits=3, max_rows=10)
    paragraph(doc, "假设包括：质量指标的方向先按指标协议统一，再用 A1 经验 CDF 标准化；配比效应是单纯形上的相对替代；跨数据集 Loss 只有在有共同验证集或配对观测时才允许标定；所有外推结论均与实测实验分开。")

    doc.add_heading("3 第一问 质量评分与冲突检验", 1)
    paragraph(doc, "每个指标先按 A1 全局 ECDF 映射到 (0,1]，再在五个语义组内取均值，主评分为五组等权均值。CRITIC 权重用于相关性和离差敏感性，熵权只作对照，长度权重只改变文档贡献口径。三套权重的完整结果见 quality_weight_comparison.csv；排序稳定性见 q_domain_rank_stability.csv。")
    add_table(doc, pd.read_csv(Q1 / "conflict_permutation_test.csv"), digits=4, max_rows=5)
    paragraph(doc, "冲突定义为同一域内至少一组位于上尾、另一组位于下尾。置换检验在域内独立打乱五组列，保持每组经验分布和域样本量，重复 1000 次。实测冲突率低于独立基准，支持“各维本质同向、冲突只属少数情形”；因此不再把冲突写成普遍现象。主口径采用 P90/P10，A1 约 9.1%，扩展集分别约 13.8% 和 11.3%。")
    add_table(doc, pd.read_csv(Q1 / "conflict_p90_p10_summary.csv"), digits=4, max_rows=5)
    add_figure(doc, Q1 / "conflict_threshold_curve.png", "图 1 冲突率的阈值敏感性")
    add_figure(doc, Q1 / "fig_domain_shift_conflict.png", "图 2 语义质量与广告分类器信号的域偏移")
    paragraph(doc, "分类器诊断将 ModernBERT、FineWeb-Edu、Qurater 聚合为 G_model，将结构统计聚合为 G_struct，并以 G_noise=1-ad_en 表示清洁方向。arxiv 与 stackexchange 的高教育/高广告共现是风险证据，支持门控或降权，不足以在没有人工真值时宣称误判率。全部 22 指标直接聚合会显著提高尾部冲突率，说明五组语义聚合是稳健性而非唯一性选择，敏感性表见 conflict_all22_sensitivity.csv。")
    paragraph(doc, "Qbar(p)=sum_i p_i Q_i* 仅是跨域场景接口。中位数插补把 14 个域的质量设为同一值，导致 Qbar 方差相对只观测域重闭合口径降低 91.4%；因此在论文结论中降级为弱证据。")
    add_table(doc, pd.read_csv(Q1 / "qbar_imputation_variance_loss.csv"), digits=4, max_rows=3)

    doc.add_heading("4 第一问 配比与逐域 Loss 模型", 1)
    paragraph(doc, "配比先行闭合，再做零值乘性替换 delta，并用 Helmert 正交基得到 z(p)。对 13 个有 Loss 的域分别拟合 L_d=beta_0d+z(p)^T beta_d、二次交互 Ridge 和 L_d=c_d+k_d exp(t_d^T p)。ILR 系数是对数对比，表示相对替代：提高某域比例必然从其他域挪出，不能理解为绝对效应。")
    model_compare = pd.read_csv(Q1 / "model_form_comparison.csv")
    add_table(doc, model_compare.groupby("model", as_index=False).agg(r2=("r2", "mean"), spearman=("spearman", "median")), digits=4, max_rows=6)
    paragraph(doc, "二次交互模型使用正则化并由交叉验证选择强度，样本只有 512 组配方，交互项只用于受约束的非加性诊断。65 条二阶差分来自独立检验集上的模型预测，不能拔高为实验观测的协同或互补机制。")
    add_figure(doc, Q1 / "loss_scatter_by_domain.png", "图 3 逐域预测值与实测 Loss")
    add_figure(doc, Q1 / "loss_residuals_by_domain.png", "图 4 逐域残差")
    paragraph(doc, "A12--A15 的 63 个配方与训练配方重合，因而只能诊断估算数据的跨尺度排序和乘性生成结构，不能作为独立实验。反演得到 gamma(10B) 与 gamma(70B) 的 pooled Pearson r=0.9997，但 gamma 与 1M Loss 的相关为 0.3591，真实 1M 到 60M 对照为 -0.1563；因此不能把人工耦合 r>0.90 写成已证实规律。")
    add_figure(doc, Q1 / "fig_est_gamma_coupling.png", "图 5 估算标度指数的耦合与真实对照")

    doc.add_heading("5 第二问 经典标度律与物理纠偏", 1)
    paragraph(doc, "B1 为 8 个 N 乘 147 个 D 的完整无重复解析网格。经典式取消独立截距：L=A N^-alpha+B D^-beta。beta 先固定为 B1 先验 0.2799，再在第二阶段以 lambda=10 的惩罚局部松弛；该流程切断 c 与 A/B 的尺度互相抵消，并报告 beta 的偏离而非虚构额外泛化能力。")
    add_table(doc, pd.read_csv(Q2 / "q2_b1_fit_coefficients.csv"), digits=4, max_rows=8)
    add_table(doc, pd.read_csv(Q2 / "q2_scaling_validation_summary.csv"), digits=4, max_rows=8)
    add_figure(doc, Q2 / "fig_q2_b1_observed_predicted.png", "图 6 B1 观测与无截距标度律预测")
    add_figure(doc, Q2 / "fig_q2_b3_trajectory_validation.png", "图 7 B3 轨迹验证")
    paragraph(doc, "B6--B8 的 Q_score 与 Loss 正相关，实测 ∂L/∂Q_score 全部为正。由于 1-Q_score 与 |Q_score-1| 只是保号仿射变换，镜像不能翻转方向。采用非对称缺陷项 L=A N^-alpha+B D^-beta+gamma Q_score^theta，gamma、theta>0，则 ∂L/∂Q_quality<0；Q_quality 趋近 1 时缺陷项趋零，自动退化为经典 N-D 标度律。")
    qtab = pd.read_csv(Q2 / "q2_quality_effect_coefficients.csv")
    cols = [c for c in ["lambda", "A", "B", "gamma", "alpha", "beta", "theta", "beta_prior_B1", "gamma_ci_low", "gamma_ci_high", "theta_ci_low", "theta_ci_high", "equivalent_sample_multiplier_ci_low", "equivalent_sample_multiplier_ci_high", "r2", "mae"] if c in qtab.columns]
    add_table(doc, qtab[cols], digits=4, max_rows=8)
    add_table(doc, pd.read_csv(Q2 / "q2_quality_mirror_equivalence.csv"), digits=4, max_rows=4)
    paragraph(doc, "lambda 敏感性中参数变化很小，说明惩罚梯度平坦，实际结果主要由 B1 先验和数据结构决定；选择模型的首要依据是物理符号正确，而不是追求 R² 提升。低端 Q_score=0.05/0.10 的配对检验显示存在与 N 相关的截断和饱和，故局部反常不解释为质量方向翻转。")
    add_figure(doc, Q2 / "fig_q2_quality_effect.png", "图 8 Q_score 缺陷率与验证 Loss")

    doc.add_heading("6 第二问 MRTS 与配比接口", 1)
    paragraph(doc, "对 L(N,D,Q,p) 做全微分并施加 sum_i dp_i=0，得到 dL=(∂L/∂lnN)dlnN+(∂L/∂lnD)dlnD+(∂L/∂Q)dQ+sum_i(∂L/∂p_i)dp_i=0。质量提升 0.1 时，参数倍数 exp(Delta ln N) 小于 1；其倒数是单靠增加参数获得同等降损所需的倍数。完整 N-D 网格的中位数和四分位区间见 MRTS 汇总，不用单点结论替代整体波动。")
    add_table(doc, pd.read_csv(Q2 / "q2_mrts_summary.csv"), digits=4, max_rows=14)
    add_figure(doc, Q2 / "fig_q2_elasticity_heatmap.png", "图 9 参数弹性、数据弹性和质量边际效用")
    paragraph(doc, "配比的工程接口写为 D_eff=D·Qbar^kappa·phi(p)。当前没有配方级 Q 与 B6--B8 Loss 的共同配对观测，因此 p 尚未进入联合可识别模型；识别需要同时变化的 N、D、p 以及同一验证集上的 Loss。问题一导出的 65 条二阶非加性只作为条件模型量。")
    add_table(doc, pd.read_csv(Q2 / "q2_second_order_nonadditivity_summary.csv"), digits=4, max_rows=4)
    bridge = pd.read_csv(Q2 / "q2_affine_bridge_test_results.csv")
    add_table(doc, bridge[[c for c in ["tier", "n", "u", "v", "r2", "F", "p_value", "ci_width_u", "identifiability"] if c in bridge.columns]], digits=4, max_rows=6)
    paragraph(doc, "跨语料假设为 L_Pile=u L_Pythia+v。High 层只有 7/75 条同模型同验证集配对，区间宽，当前结论是已提出可检验假设但定量识别不足；B9/B10 只用于超百亿外推的情景敏感性。")

    doc.add_heading("7 模型评价与推广", 1)
    paragraph(doc, "评价时严格分开同尺度绝对 Loss、跨尺度绝对误差和排序相关。问题一的 1M 到 1M 测试宏平均 R² 约 0.7641，跨到 60M 与 1B 时绝对 R² 显著恶化而秩相关仍较高，说明绝对 Loss 不宜跨尺度迁移、相对排序更稳定。问题二的 B1 网格结果是解析识别，不作为外部泛化承诺。")
    paragraph(doc, "推广时建议将 Qbar 作为弱先验或门控变量，先在有配对观测的子集上估计 kappa 和 phi(p)，再评估其对逐域 Loss 的增量解释。对于新领域，优先报告域内标准化后的排序、替代效应和不确定性，避免把中位数插补的宏观质量标量误读为实验测量。")
    add_figure(doc, Q1 / "q_domain_boxplot.png", "图 10 17 域质量分布和不确定性")

    doc.add_heading("参考文献", 1)
    paragraph(doc, "[1] Aitchison J. The Statistical Analysis of Compositional Data. Chapman and Hall, 1986.")
    paragraph(doc, "[2] Egozcue J J et al. Isometric logratio transformations for compositional data analysis. Mathematical Geology, 2003.")
    paragraph(doc, "[3] Kaplan J et al. Scaling Laws for Neural Language Models. arXiv:2001.08361, 2020.")
    paragraph(doc, "[4] Hoffmann J et al. Training Compute-Optimal Large Language Models. arXiv:2203.15556, 2022.")
    paragraph(doc, "[5] Data Mixing Laws and RegMix related studies, with model comparison restricted to the supplied A and B attachments.")

    DOC.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOC)
    print(DOC)


if __name__ == "__main__":
    main()
