"""Build the single complete Question 2 report from generated outputs."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs" / "q2"
DOC = OUT / "第二问_完整报告.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table):
    tblPr = table._tbl.tblPr; b = tblPr.first_child_found_in("w:tblBorders")
    if b is None: b = OxmlElement("w:tblBorders"); tblPr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = b.find(qn("w:" + edge))
        if el is None: el = OxmlElement("w:" + edge); b.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "D9D9D9")

def table(doc, frame, digits=4, max_rows=18):
    f = frame.head(max_rows).copy()
    t = doc.add_table(rows=1, cols=len(f.columns)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"; borders(t)
    for j, c in enumerate(f.columns):
        cell=t.rows[0].cells[j]; cell.text=str(c); shade(cell,"1F4E78"); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cell.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
    for i, row in f.iterrows():
        cells=t.add_row().cells
        for j, c in enumerate(f.columns):
            v=row[c]
            if isinstance(v,float): s=f"{v:.{digits}f}"
            else: s=str(v)
            cells[j].text=s; cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                for r in p.runs: r.font.size=Pt(8)
        if len(t.rows)%2==0:
            for cell in cells: shade(cell,"F3F6F9")
    return t

def para(doc, text, bold_lead=None):
    p=doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.15
    return p

def fig(doc, filename, caption, width=6.2):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(OUT/filename), width=Inches(width))
    c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic=True; c.runs[0].font.size=Pt(9)

def main():
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    styles=d.styles; styles["Normal"].font.name="Microsoft YaHei"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); styles["Normal"].font.size=Pt(10)
    for s in ["Title","Heading 1","Heading 2"]:
        styles[s].font.name="Microsoft YaHei"; styles[s]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); styles[s].font.color.rgb=RGBColor(0,0,0)
    title_style = styles["Title"]; style_ppr = title_style._element.pPr
    if style_ppr is not None:
        style_bdr = style_ppr.find(qn("w:pBdr"))
        if style_bdr is not None: style_ppr.remove(style_bdr)
    title=d.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run("问题二 跨维度数据融合与广义标度律").bold=True
    # Word's built-in Title style can carry a theme border; remove it for a clean report title.
    ppr = title._p.get_or_add_pPr(); pbd = ppr.find(qn("w:pBdr"))
    if pbd is not None: ppr.remove(pbd)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("基于 B1–B10 附件与问题一接口的完整建模报告").italic=True
    para(d,"本问建立参数规模、数据量、质量与语料配比共同作用下的验证损失模型，并用跨模型族、轨迹、文献和大规模估算数据检验其可迁移边界。报告中的绝对 Loss 只在 B 附件的 Pythia/Cerebras 验证集口径内比较；问题一的 The Pile Loss 不与 B 的 val_loss 直接拼接。")
    d.add_heading("1 结论摘要",1)
    para(d,"B1 的 Pythia 主拟合采用 L=c+A N^{-alpha}+B D^{-beta}，估计 alpha=0.33996、beta=0.27988。随机留出测试 R²=1.0000、MAE=0.0001，五折 CV 的 R² 均为 1.0000，说明 B1 轨迹由稳定的幂律结构生成，参数估计具有很小的抽样波动。")
    para(d,"跨数据源结果分成两类：B4 跨族收敛点 R²=0.6045、秩相关 0.9830；B5 文献点 R²=0.7306、秩相关 0.9588，说明尺度排序可迁移。B2 Cerebras 的绝对 Loss 受验证集和模型族偏移影响，R²=-5.0728，但秩相关仍为 0.7881，因此不把 B1 的绝对截距直接外推到 B2。")
    para(d,"B6–B8 使用其原生半合成 Q_score 单独拟合质量修正，R²=0.5912；Q 的边际效应随 N、D 变化。问题一接口中的 Q_star 与 ILR 配比只用于场景接口和域替代解释，未被伪装成 B6–B8 的配对监督样本。")
    d.add_heading("2 数据与口径边界",1)
    boundary=pd.read_csv(OUT/"q2_input_usage_boundary.csv"); table(d,boundary, max_rows=10)
    para(d,"B1 是真实 Pythia 训练日志，作为经典 N-D 主拟合；B2 是 Cerebras 轨迹，用于模型族外验证；B3 是插值轨迹，用于同族轨迹验证；B4/B5 分别用于跨族和文献验证；B6–B8 标记为半合成或估算，只用于识别原生 Q_score 的质量效应；B9/B10 是 100B 以上元数据与估算 Loss，只用于情景外推。A 与 B 的验证集不同，第一问接口中的 Q_star 不是 B6–B8 的 Q_score 量表。")
    d.add_heading("3 标度律模型",1)
    para(d,"对 B1 的每个观测点使用 N_params_B（单位 B）和 D_tokens_B（单位 B tokens）。主模型为：")
    para(d,"L(N,D)=c+A N^{-alpha}+B D^{-beta}。其中 c 是不可由本次扩展继续降低的残差基线，A、B 是两类规模项的幅度，alpha、beta 是参数量与数据量的弹性衰减指数。这个形式在 N 或 D 固定时退化为单变量幂律，在二者同时增加时给出可解释的边际收益。")
    coeff=pd.read_csv(OUT/"q2_b1_fit_coefficients.csv"); table(d,coeff, digits=6, max_rows=10)
    cv=pd.read_csv(OUT/"q2_b1_cv_metrics.csv"); table(d,cv,digits=6,max_rows=8); fig(d,"fig_q2_b1_observed_predicted.png","图 1 B1 观测值与经典标度律预测值")
    fig(d,"fig_q2_b1_residuals.png","图 2 B1 残差随参数规模的诊断")
    d.add_heading("4 外部验证",1)
    ext=pd.read_csv(OUT/"q2_scaling_validation_summary.csv"); table(d,ext,digits=4,max_rows=10)
    traj=pd.read_csv(OUT/"q2_b3_trajectory_validation.csv"); table(d,traj[["trajectory","n","r2","mae","spearman"]],digits=5,max_rows=10); fig(d,"fig_q2_b3_trajectory_validation.png","图 3 B3 八条轨迹验证，实线为观测、虚线为预测")
    para(d,"B3 的八条轨迹 R² 均约 0.99995，说明同一 Pythia 族内的插值路径与主模型一致。B4/B5 的绝对误差较大但秩相关高，表明跨族使用时应优先采用排序、相对改善率或重新估计截距，而不能把 B1 的绝对 Loss 当作所有模型族的共同标尺。")
    d.add_heading("5 质量扩展",1)
    para(d,"B6–B8 内部采用固定的 alpha、beta，把质量作用写成带交互的扩展：L_Q=b0+b_N N^{-alpha}+b_D D^{-beta}+b_Q(Q-Q0)+b_NQ(Q-Q0)N^{-alpha}+b_DQ(Q-Q0)D^{-beta}。Q0 为原生 Q_score 中位数 0.55，岭回归正则强度通过五折 CV 选择。该模型只解释 B6–B8 的原生 Q_score，不能直接替代问题一的 A1 ECDF Q_star。")
    qc=pd.read_csv(OUT/"q2_quality_effect_coefficients.csv"); table(d,qc,digits=6,max_rows=10); fig(d,"fig_q2_quality_effect.png","图 4 B6–B8 原生 Q_score 与验证 Loss")
    elast=pd.read_csv(OUT/"q2_elasticities.csv"); key=elast[(elast.Q_score==.5)&(elast.N_params_B.isin([.07,1,10,100,700]))&(elast.D_tokens_B.isin([10,100,1000,10000]))]; table(d,key[["N_params_B","D_tokens_B","Q_score","loss","elasticity_N","elasticity_D","marginal_quality"]],digits=5,max_rows=25); fig(d,"fig_q2_elasticity_heatmap.png","图 5 Q=0.5 时 Loss 对参数规模的弹性")
    para(d,"弹性为 ∂L/∂ln N 除以 L 或 ∂L/∂ln D 除以 L；数值越负，继续增加该资源带来的相对 Loss 降幅越大。质量边际项为 ∂L/∂Q，在当前半合成口径下可与参数和数据边际收益比较，但其外推必须标记为情景分析。")
    d.add_heading("6 大规模外推",1)
    large_m=pd.read_csv(OUT/"q2_large_scale_metrics.csv"); table(d,large_m,digits=5,max_rows=5); fig(d,"fig_q2_large_scale_extrapolation.png","图 6 B9/B10 大规模估算点与经典模型外推")
    para(d,"B9/B10 的 R²=0.99999 只说明该估算表与 B1 形式高度一致，不构成独立实验验证。报告将其作为资源配置情景工具：给定 N、D 和质量水平，可以计算相对 Loss 变化，并同时保留估算数据的边界标签。")
    d.add_heading("7 问题一接口与配比替代",1)
    para(d,"问题一输出的 q2_input_interface_bundle.json 提供 17 域 Q_star、训练配方基线 Q0、ILR/替代向量和二阶组合对。本问不重新读取 A 原始质量记录，而是将其作为外部场景输入。宏观质量泛函为 Q_bar(p)=sum_i p_i Q_i_star；配比变量仍按问题一的单纯形闭合、乘性零值替换和 16 维 ILR 语义解释，系数表示相对替代而非绝对增加。")
    dom=pd.read_csv(OUT/"q2_domain_substitution_effects.csv"); table(d,dom,digits=5,max_rows=20)
    para(d,"Q_star_status 为 observed 的域来自问题一文档级观测；其余域为问题一接口中的中位数填充，因此配比预测的不确定性主要来自这些未观测域。若第二问需要把 Q_bar(p) 代入质量扩展模型，应先做单调重标定，再将重标定后的质量值作为情景变量，不能把 A1 Q_star 直接当成 B6–B8 Q_score。")
    d.add_heading("8 资源配置解释",1)
    para(d,"模型给出三条可操作结论。第一，增加 N 或 D 的收益具有递减性，alpha 和 beta 分别控制递减速度；在固定算力约束下，应比较两者的单位成本边际下降，而不是只追求参数规模。第二，质量提升项与 N、D 交互，意味着低质量数据不能简单用更多参数补偿，质量门控适合作为先验筛选或情景约束。第三，配比影响使用 ILR 替代语义：提高一个域的比例必然从其他域挪出，二阶协同或拮抗应在固定总配比下比较。")
    d.add_heading("9 复现说明",1)
    para(d,"在项目根目录执行 `$env:PYTHONPATH='src'; .venv\\Scripts\\python.exe -m f_question.q2` 可重新生成 outputs/q2 下的 CSV 和 PNG。报告由 build_q2_docx.py 根据这些结果生成。原始数据集保留在本地项目中，公开仓库只提交代码、报告和小型结果产物。")
    d.save(DOC); print(DOC)

if __name__ == "__main__": main()
