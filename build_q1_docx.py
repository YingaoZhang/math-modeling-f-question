from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\project")
OUT = ROOT / "outputs" / "q1"
MAIN = OUT / "第一问_完整建模报告.md"
SUPP = OUT / "第一问_补充实证与解释.md"
DOCX = OUT / "第一问_完整报告.docx"

FIGURES = {
    "质量分域分布箱线图": "q_domain_boxplot.png",
    "github 与 arxiv 质量分比较": "q_github_arxiv_comparison.png",
    "冲突阈值敏感性曲线": "conflict_threshold_curve.png",
    "分类器域偏移与冲突": "fig_domain_shift_conflict.png",
    "估算数据标度指数耦合": "fig_est_gamma_coupling.png",
    "逐域 Loss 预测散点": "loss_scatter_by_domain.png",
    "逐域 Loss 残差": "loss_residuals_by_domain.png",
    "标准化跨尺度排序诊断": "loss_scatter_standardized_aggregate.png",
}

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', sz='4'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), sz)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def set_cell_text(cell, text, bold=False, size=8.2):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '等线')
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for j in range(ncols):
        set_cell_text(hdr[j], rows[0][j] if j < len(rows[0]) else '', bold=True, size=7.8)
        set_cell_shading(hdr[j], '1F4E78')
        for run in hdr[j].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255,255,255)
        set_cell_border(hdr[j])
    for i, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for j in range(ncols):
            val = row[j] if j < len(row) else ''
            set_cell_text(cells[j], val, size=7.4)
            if i % 2 == 1:
                set_cell_shading(cells[j], 'F2F6FA')
            set_cell_border(cells[j])
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def split_md_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith('|'):
        line = lines[i].strip().strip('|')
        vals = [v.strip() for v in line.split('|')]
        if not all(re.fullmatch(r':?-{2,}:?', v) for v in vals):
            rows.append(vals)
        i += 1
    return rows, i

def add_image(doc, filename, caption):
    path = OUT / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(6.25))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(90,90,90)

def clean_inline(s):
    s = s.replace('**', '').replace('`', '')
    s = re.sub(r'\$\$(.*?)\$\$', r'\1', s)
    s = re.sub(r'\$(.*?)\$', r'\1', s)
    s = s.replace('\\mathbf', '').replace('\\operatorname', '')
    return s

def add_markdown(doc, path, include_heading=True):
    lines = path.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith('|') and i + 1 < len(lines) and '---' in lines[i+1]:
            rows, i = split_md_table(lines, i)
            add_table(doc, rows)
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            txt = clean_inline(m.group(2))
            if level == 1:
                if include_heading:
                    p = doc.add_paragraph(style='Title')
                    p.add_run(txt)
            else:
                p = doc.add_heading(txt, level=min(level, 3))
                p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if line.startswith('---'):
            i += 1
            continue
        if line.startswith('>'):
            line = line[1:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        text = clean_inline(line)
        p.add_run(text)
        i += 1

def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.7); sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '等线')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(20,20,20)
    for name, size in [('Title',20),('Heading 1',15),('Heading 2',12.5),('Heading 3',11)]:
        st = styles[name]
        st.font.name = 'Aptos Display'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), '等线')
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0,0,0)
        st.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
        st.paragraph_format.space_after = Pt(5)

doc = Document()
configure(doc)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('数学建模 F 题第一问\n语料质量评分与配比到验证损失建模完整报告')
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run('数据版本：A1–A18；主分析输出目录：outputs/q1').italic = True
doc.add_paragraph()

doc.add_heading('报告摘要', level=1)
doc.add_paragraph('本报告将第一问的质量评分、冲突识别、配比建模、外推诊断和跨问接口统一在一份文档中。质量分使用 A1 全局经验 CDF 映射，始终满足 Q∈(0,1]；配比先闭合到单纯形，再做乘性零值替换和 16 维 Helmert ILR；13 个 The Pile 验证域分别建模，未对原始 Loss 做算术平均。')
doc.add_paragraph('主要结论：同尺度 1M→1M 的逐域建模有效，独立检验宏平均 R² 约 0.889（指数模型），而绝对 Loss 无法跨尺度迁移；秩在跨尺度上仍较稳定。质量分的冲突惩罚是无监督的定义型敏感性评分，因没有文档级 Q 与配方级 Loss 配对，不能宣称其带来监督增益。A12–A15 的估算数据表现出高度规则的乘性标度结构，应作为估算机制诊断而非实验结果。')

doc.add_heading('一 研究目标与总体流程', level=1)
doc.add_paragraph('第一问包含三个相互连接但口径不同的任务：其一，从 22 个质量信号得到文档级 Q 和域级质量概况；其二，识别不同质量维度之间的冲突并构造可解释的惩罚评分；其三，在配比数据上使用成分数据分析，把 17 域配比映射为各验证域 Loss。完整流程为：方向统一 → A1 ECDF 标定 → 组内聚合与赋权 → 域内冲突诊断 → 配比闭合与 ILR → 逐 Loss 域训练/CV/独立检验 → 跨尺度和估算机制诊断 → 导出第二问接口。')

doc.add_heading('二 质量分定义与结果', level=1)
add_markdown(doc, MAIN, include_heading=False)

for caption, fname in [
    ('图 1  七个抽样域的文档质量分分布。', 'q_domain_boxplot.png'),
    ('图 2  github 与 arxiv 的质量分对比，显示明显的域间水平差异。', 'q_github_arxiv_comparison.png'),
]:
    add_image(doc, fname, caption)

doc.add_heading('三 补充实证与机制解释', level=1)
doc.add_paragraph('下面汇总补充报告中的估算数据、冲突机制、分类器域偏移、组合效应和 A17/A18 证据。数值表保留原始分析口径，所有估算结果均明确标注为估算或描述性诊断。')
add_markdown(doc, SUPP, include_heading=False)

for caption, fname in [
    ('图 3  冲突阈值敏感性曲线。阈值越严格，冲突率越低；主口径采用域内 P75/P25。', 'conflict_threshold_curve.png'),
    ('图 4  arxiv 与 stackexchange 的分类器域偏移诊断。', 'fig_domain_shift_conflict.png'),
    ('图 5  A12–A15 估算标度指数 gamma 的耦合关系。', 'fig_est_gamma_coupling.png'),
    ('图 6  各验证域独立 Loss 的预测值与实测值。', 'loss_scatter_by_domain.png'),
    ('图 7  各验证域独立模型残差。', 'loss_residuals_by_domain.png'),
    ('图 8  域内标准化后的跨尺度排序诊断。', 'loss_scatter_standardized_aggregate.png'),
]:
    add_image(doc, fname, caption)

doc.add_heading('四 结果图与数据文件索引', level=1)
doc.add_paragraph('图 1–8 已嵌入本报告。为便于复核，原始表格和模型系数仍以 CSV/JSON 保存在 outputs/q1，不改变现有 44 个产物。关键文件包括：q_domain_summary.csv、q_domain_summary_A1_A2_A3_all.csv、quality_weight_comparison.csv、conflict_threshold_by_population.csv、conflict_pair_decomposition.csv、conflict_penalty_sensitivity.csv、model_form_comparison.csv、loss_metrics_by_domain_and_scale.csv、data_mixing_law_coefficients.csv、mixture_joint_transfer_effects.csv、scaling_est_reverse_engineering.csv、conflict_classifier_domain_shift.csv、q2_input_interface_bundle.json。')
doc.add_heading('五 最终结论与边界', level=1)
doc.add_paragraph('质量分应理解为冻结在 A1 经验分布上的相对质量刻度，不能直接与附件 B 的 Q_score 混用；跨附件 Loss 也必须经过桥接，不能把 The Pile 验证 Loss 与 Pythia 自有 val_loss 当作同一量尺。配比系数表示相对替代：提高一个域的比例必然从其他域挪出份额。四个没有直接 Loss 列的域只通过闭合和替代关系间接进入系数，不能解读为它们自身的直接预测效果。A12–A15 共用同一批 63 个配方，A2/A3 与 A1 存在完全重合的子抽样关系，均不应当被当作独立外部实验。')
doc.add_paragraph('在独立同尺度检验上，指数模型优于线性基线；二次交互模型提供组合非加性诊断，但组合效应属于模型条件关联而非因果协同。跨尺度检验表明绝对 Loss 不可迁移，而域内排序仍保留，因此第二问应采用明确的桥接映射和尺度分层报告。')

doc.save(DOCX)
print(DOCX)
